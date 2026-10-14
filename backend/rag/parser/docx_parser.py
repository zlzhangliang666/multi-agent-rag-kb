"""
Word 解析器 - 使用 python-docx
"""
import logging
import os

from docx import Document as DocxDocument
from . import ParsedDocument

logger = logging.getLogger(__name__)


class DocxParser:
    HEADING_PREFIX = "heading"

    def parse(self, file_path: str, doc_id: str) -> ParsedDocument:
        file_path = os.path.abspath(file_path)
        logger.debug("Opening DOCX: %s", file_path)
        docx_doc = DocxDocument(file_path)
        text_parts = []
        for para in docx_doc.paragraphs:
            style_name = (para.style.name or "").lower() if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            if style_name.startswith(self.HEADING_PREFIX):
                level = 0
                for ch in style_name.replace(self.HEADING_PREFIX, "").strip():
                    if ch.isdigit():
                        level = int(ch); break
                marker = "#" * max(level, 1)
                text_parts.append(f"{marker} {text}")
            else:
                text_parts.append(text)
        full_text = "\n".join(text_parts)
        metadata = {}
        try:
            cp = docx_doc.core_properties
            if cp:
                if cp.author: metadata["author"] = cp.author
                if cp.title: metadata["title"] = cp.title
        except Exception:
            pass
        return ParsedDocument(doc_id=doc_id, doc_name=os.path.basename(file_path),
                              pages=[{"page_num": 1, "text": full_text}],
                              metadata=metadata)